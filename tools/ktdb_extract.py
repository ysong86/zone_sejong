# -*- coding: utf-8 -*-
"""KTDB 대전세종충청권 여객 OD(목적·주수단) 엑셀에서 세종 관련 행만 뽑는다.

  python tools/ktdb_extract.py [연도시트=2024년]
    입력: data/raw/ktdb/obj_od.xlsx, mod_od.xlsx, obj_zone.xlsx(존 체계)
    출력: data/raw/ktdb/sejong_od_<연도>.json

원 엑셀은 각 250MB·46만 행이라 매번 읽기엔 무겁다. 세종 존(권역 존체계 읍면동)이 출발이나
도착인 행과, 존 번호 → 시도·시군구·행정동 표만 남긴다.
"""
from __future__ import annotations

import json
import os
import sys
import time

import openpyxl

HERE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DIR = os.path.join(HERE, "data", "raw", "ktdb")


def zones():
    wb = openpyxl.load_workbook(os.path.join(DIR, "obj_zone.xlsx"), read_only=True)
    ws = wb["대전세종충청권"]
    out = {}
    for i, r in enumerate(ws.iter_rows(values_only=True)):
        if i < 2:
            continue
        sido, sgg, dong, zno = r[6], r[7], r[8], r[9]      # 2025년 사업 존체계
        if isinstance(zno, (int, float)):
            out[int(zno)] = {"sido": sido, "sgg": sgg, "dong": dong, "inner": r[11] == 1}
    return out


def extract(path, sheet, keep):
    t = time.time()
    wb = openpyxl.load_workbook(path, read_only=True)
    ws = wb[sheet]
    head, rows = None, []
    for i, r in enumerate(ws.iter_rows(values_only=True)):
        if i == 0:
            head = [c for c in r if c is not None]
            continue
        o, d = r[0], r[1]
        if o in keep or d in keep:
            rows.append([int(o), int(d)] + [round(float(x or 0), 2) for x in r[2:len(head)]])
        if i % 100000 == 0:
            print(f"  {os.path.basename(path)} {i:,} 행 · {time.time() - t:.0f}s", flush=True)
    return head, rows


def main(sheet="2024년"):
    zt = zones()
    sj = {z for z, v in zt.items() if (v["sido"] or "").startswith("세종")}
    print("세종 존", sorted(sj))
    oh, orows = extract(os.path.join(DIR, "obj_od.xlsx"), sheet, sj)
    mh, mrows = extract(os.path.join(DIR, "mod_od.xlsx"), sheet, sj)
    out = {"sheet": sheet, "zones": {str(k): v for k, v in zt.items()}, "sejong": sorted(sj),
           "obj_head": oh, "obj": orows, "mod_head": mh, "mod": mrows}
    path = os.path.join(DIR, f"sejong_od_{sheet.rstrip('년')}.json")
    with open(path, "w", encoding="utf-8") as f:
        json.dump(out, f, ensure_ascii=False)
    print("저장:", path, len(orows), len(mrows), "행", mh)


if __name__ == "__main__":
    main(*(sys.argv[1:2]))
